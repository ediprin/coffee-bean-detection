from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path
import re, argparse

GREEN = RGBColor(77, 166, 62)
DARK_GREEN = RGBColor(45, 130, 44)
PALE_GREEN = RGBColor(229, 242, 220)
PALE_BLUE = RGBColor(202, 236, 246)
DARK = RGBColor(25,25,25)
WHITE = RGBColor(255,255,255)
GRAY = RGBColor(90,90,90)


def parse_source(path):
    text=Path(path).read_text(encoding='utf-8')
    blocks=re.split(r'(?m)^---\s*$',text); out=[]
    for b in blocks:
        m=re.search(r'(?m)^## SLIDE\s+(\d+)\s+—\s+(.+?)\s*$',b)
        if not m: continue
        body=b[m.end():].strip(); tm=re.search(r'(?m)^###\s+(.+?)\s*$',body)
        title=tm.group(1).strip() if tm else ''
        if tm: body=body[tm.end():].strip()
        out.append(dict(number=int(m.group(1)),label=m.group(2).strip(),title=title,body=body))
    return sorted(out,key=lambda x:x['number'])

def clean(s):
    s=re.sub(r'```text\s*','',s); s=s.replace('```','')
    s=re.sub(r'\*\*(.*?)\*\*',r'\1',s); s=re.sub(r'\*(.*?)\*',r'\1',s)
    s=s.replace('\\*','*')
    return s.strip()

def delete_all_slides(prs):
    while len(prs.slides):
        slide_id=prs.slides._sldIdLst[0]; prs.part.drop_rel(slide_id.rId); del prs.slides._sldIdLst[0]

def add_title(slide, text, size=24):
    box=slide.shapes.add_textbox(Inches(.9),Inches(.75),Inches(10.4),Inches(.6))
    p=box.text_frame.paragraphs[0]; p.text=text; p.font.size=Pt(size); p.font.bold=False; p.font.color.rgb=DARK
    return box

def add_text(slide, text, x,y,w,h, size=17, bold=False, color=DARK, align=PP_ALIGN.LEFT, margin=.08, valign=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=Inches(margin); tf.vertical_anchor=valign
    lines=text.split('\n')
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.alignment=align
        for r in p.runs: r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box

def add_panel(slide,x,y,w,h,fill=PALE_GREEN,line=None,radius=True):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb=fill
    shp.line.color.rgb = line if line else fill
    return shp

def add_number_item(slide,num,text,y,x=.95,text_w=4.0,size=16):
    shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(.6),Inches(.42)); shp.fill.solid(); shp.fill.fore_color.rgb=GREEN; shp.line.color.rgb=GREEN
    add_text(slide,f'{num:02d}',x+.07,y+.01,.46,.3,size=15,bold=True,color=WHITE,align=PP_ALIGN.CENTER)
    add_text(slide,text,x+.77,y-.02,text_w,.72,size=size)

def slide_cover(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[0])
    for sh in s.shapes:
        if getattr(sh,'has_text_frame',False): sh.text_frame.clear()
    title=s.shapes.title; title.text=d['title'];
    for p in title.text_frame.paragraphs:
        for r in p.runs: r.font.size=Pt(24); r.font.bold=True; r.font.color.rgb=WHITE
    sub=s.placeholders[1]; sub.text='[Nama Mahasiswa]\n[NIM]'
    for p in sub.text_frame.paragraphs:
        for r in p.runs: r.font.size=Pt(12); r.font.color.rgb=WHITE
    add_text(s,'Seminar Proposal',2.1,1.1,3.2,.55,size=17,bold=True,color=WHITE)
    add_text(s,'Dosen Pembimbing\n[Dosen Pembimbing 1]\n[Dosen Pembimbing 2]',7.55,5.42,4.6,.7,size=10,color=WHITE)

def slide_agenda(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Pembahasan Materi',26)
    add_number_item(s,1,'Pendahuluan',2.65,x=.95,text_w=3.8)
    add_number_item(s,2,'Penelitian Terdahulu',2.65,x=6.65,text_w=4.0)
    add_number_item(s,3,'Metodologi Penelitian',4.05,x=.95,text_w=4.4)

def slide_intro(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Pendahuluan',24)
    pts=[
    'Inspeksi mutu biji kopi hijau masih banyak bergantung pada pengamatan visual, sehingga konsistensinya dapat dipengaruhi pengalaman dan kondisi pemeriksa.',
    'Deteksi otomatis menjadi lebih menantang ketika kategori cacat semakin rinci karena beberapa kelas memiliki kemiripan pada warna, tekstur, bentuk, dan detail lokal.',
    'Kondisi tersebut mendorong kebutuhan representasi citra yang lebih diskriminatif untuk deteksi fine-grained cacat biji kopi.' ]
    ys=[2.0,3.25,4.5]
    for i,(t,y) in enumerate(zip(pts,ys),1): add_number_item(s,i,t,y,text_w=9.8,size=17)

def slide_problem(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Rumusan Masalah',23)
    add_panel(s,.45,1.55,11.9,3.55,fill=RGBColor(180,215,168),radius=False)
    txt='Deteksi fine-grained cacat biji kopi menghadapi kemiripan visual antarkelas, sedangkan pemanfaatan informasi frekuensi-angular sebelum proses deteksi masih terbatas. Penelitian ini mengkaji penerapan dan optimasinya pada YOLO26n terhadap kinerja deteksi dan biaya komputasi.'
    add_text(s,txt,.9,2.0,10.9,2.35,size=18,align=PP_ALIGN.JUSTIFY,valign=MSO_ANCHOR.MIDDLE)

def slide_scope(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'BATASAN MASALAH',22)
    add_panel(s,.35,1.5,12.0,4.9,fill=RGBColor(221,239,204),radius=False)
    items=['Deteksi fine-grained cacat biji kopi hijau.','Dataset utama robusta_SNI_Dataset (21 kelas).','Dataset Capstone, Lulus, dan Niacubilla sebagai konfirmasi.','Model utama YOLO26n tanpa modifikasi backbone, neck, dan head.','Optimasi difokuskan pada prapemrosesan frekuensi-angular.','Evaluasi utama menggunakan mAP50–95 dan biaya komputasi end-to-end.']
    y=1.85
    for t in items:
        add_text(s,'• '+t,.65,y,11.3,.52,size=15); y+=.7

def slide_related(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Penelitian Terdahulu',22)
    rows=[('Hong et al. (2026)','Improved YOLOv10','Deteksi cacat kopi'),('Jiao et al. (2025)','Multistage fusion + attention','Diskriminasi fitur cacat kopi'),('Li et al. (2025)','Fourier preprocessing + YOLO','Pemrosesan spektral sebelum deteksi'),('Xu et al. (2025)','AFAB','Frekuensi-angular untuk fine-grained detection')]
    x=[.55,3.2,7.45]; widths=[2.6,4.2,4.55]; y=1.55; rh=.72
    headers=['Penelitian','Pendekatan','Fokus']
    for xi,w,h in zip(x,widths,headers):
        add_panel(s,xi,y,w,rh,fill=GREEN,radius=False); add_text(s,h,xi+.06,y+.09,w-.12,.45,size=14,bold=True,color=WHITE,align=PP_ALIGN.CENTER)
    y+=rh
    for r,row in enumerate(rows):
        fill=RGBColor(239,246,235) if r%2==0 else RGBColor(224,237,218)
        for xi,w,txt in zip(x,widths,row):
            add_panel(s,xi,y,w,rh,fill=fill,line=RGBColor(190,210,182),radius=False); add_text(s,txt,xi+.08,y+.1,w-.16,.48,size=11.5,align=PP_ALIGN.CENTER)
        y+=rh
    add_panel(s,.55,5.05,11.45,.88,fill=RGBColor(217,235,203),radius=False)
    add_text(s,'Gap: penelitian pada kopi lebih banyak mengembangkan representasi internal model; prapemrosesan frekuensi-angular sebelum detektor belum dikaji secara khusus.',.72,5.18,11.05,.58,size=12.5,bold=True)

def slide_why(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Mengapa Frekuensi-Angular?',22)
    cards=[('Fine-grained defect','Perbedaan kecil pada tekstur dan pola permukaan'),('Frequency','Menangkap karakteristik perubahan dan detail visual'),('Angular','Menangkap distribusi respons berdasarkan arah')]
    xs=[.7,4.45,8.2]
    for x,(h,b) in zip(xs,cards):
        add_panel(s,x,2.0,3.15,2.45,fill=RGBColor(223,239,213))
        add_text(s,h,x+.15,2.25,2.85,.45,size=17,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
        add_text(s,b,x+.25,3.0,2.65,1.0,size=14,align=PP_ALIGN.CENTER)
    add_text(s,'Hipotesis: representasi frekuensi-angular dapat membantu menghasilkan masukan yang lebih diskriminatif bagi detektor.',1.0,5.05,10.8,.72,size=15,bold=True,align=PP_ALIGN.CENTER)

def slide_dataset(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Dataset Penelitian',22)
    add_panel(s,.75,1.7,5.2,3.45,fill=RGBColor(227,241,217)); add_panel(s,6.4,1.7,5.2,3.45,fill=RGBColor(227,241,217))
    add_text(s,'Dataset Utama',1.0,2.0,4.7,.45,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'robusta_SNI_Dataset\n21 kelas\n\nPengembangan dan pemilihan C*',1.0,2.65,4.7,1.75,size=16,align=PP_ALIGN.CENTER)
    add_text(s,'Dataset Konfirmasi',6.65,2.0,4.7,.45,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'Capstone — 14 kelas\nLulus — 6 kelas\nNiacubilla — 9 kelas',6.7,2.75,4.6,1.55,size=16,align=PP_ALIGN.CENTER)
    add_text(s,'Split 70% train · 15% validation · 15% test  |  Setiap dataset digunakan secara terpisah.',1.2,5.55,10.4,.52,size=13,bold=True,align=PP_ALIGN.CENTER)

def slide_method(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'METODOLOGI PENELITIAN',22)
    add_text(s,'Empat kondisi eksperimen utama:',.9,1.65,5.0,.45,size=15)
    labs=[('B0','YOLO26n'),('B1','CLAHE → YOLO26n'),('B2','C0 → YOLO26n'),('B3','C* → YOLO26n')]
    xs=[.75,3.75,6.75,9.75]
    for x,(a,b) in zip(xs,labs):
        add_panel(s,x,2.25,2.3,1.15,fill=RGBColor(209,232,193)); add_text(s,a,x+.08,2.4,.55,.55,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER); add_text(s,b,x+.65,2.42,1.55,.55,size=13,align=PP_ALIGN.CENTER)
    add_panel(s,.9,4.0,11.2,1.65,fill=RGBColor(221,239,204),radius=False)
    add_text(s,'B2 − B0  → efek frequency-angular reference frontend\nB3 − B2  → efek optimasi desain\nB3 − B1  → perbandingan terhadap CLAHE',1.25,4.25,10.5,1.15,size=14,align=PP_ALIGN.CENTER)

def flow_slide(prs,title,steps,footer=None):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,title,22)
    n=len(steps); y0=1.6; gap=.12; avail=4.7; h=(avail-gap*(n-1))/n
    for i,step in enumerate(steps):
        y=y0+i*(h+gap); add_panel(s,3.05,y,6.7,h,fill=RGBColor(226,240,216)); add_text(s,step,3.25,y+.05,6.3,h-.1,size=12.5,bold=(i in [0,n-1]),align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
        if i<n-1: add_text(s,'↓',6.0,y+h-.03,.75,.25,size=15,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    if footer: add_text(s,footer,1.0,6.15,10.9,.4,size=12.5,bold=True,align=PP_ALIGN.CENTER)
    return s

def slide_preprocess(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1])
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Alur Prapemrosesan Frekuensi-Angular',22)
    steps=['Citra RGB','Patch Lokal','FFT 2D','Analisis\nAmplitudo & Arah','Adaptive Spectral\nWeighting','Inverse FFT','Rekonstruksi','Residual Fusion','YOLO26n']
    xs=[.65,4.5,8.35]; ys=[1.75,3.15,4.55]
    k=0
    for r,y in enumerate(ys):
        order=range(3) if r%2==0 else range(2,-1,-1)
        for c in order:
            x=xs[c]; step=steps[k]; k+=1
            fill=RGBColor(209,232,193) if step in ('Citra RGB','YOLO26n') else RGBColor(229,242,220)
            add_panel(s,x,y,3.0,.9,fill=fill)
            add_text(s,step,x+.12,y+.12,2.76,.62,size=14.5,bold=step in ('Citra RGB','YOLO26n'),align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
            if k<9:
                if (r%2==0 and c<2) or (r%2==1 and c>0):
                    ax=x+3.03 if r%2==0 else x-.48
                    add_text(s,'→' if r%2==0 else '←',ax,y+.2,.45,.4,size=20,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
        if r<2:
            dropx=11.55 if r==0 else .7
            add_text(s,'↓',dropx,y+.95,.45,.36,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'I′ = I + I ⊙ G   ·   Parameter-free frontend   ·   YOLO26n tidak dimodifikasi',1.0,5.95,10.8,.45,size=13.5,bold=True,align=PP_ALIGN.CENTER)


def slide_design(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1])
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Optimasi Desain',22)
    stages=[('C0','Reference\nfrequency-angular'),('C1','+ Hann window'),('C2','+ Unsigned\norientation'),('C3','+ Radial bands'),('C4','+ Soft threshold'),('C5','+ Luminance\nguidance'),('C*','Konfigurasi\nterpilih')]
    x0=.45; y=2.15; w=1.55; gap=.22
    for i,(code,desc) in enumerate(stages):
        x=x0+i*(w+gap)
        fill=RGBColor(203,229,190) if code in ('C0','C*') else RGBColor(229,242,220)
        add_panel(s,x,y,w,2.25,fill=fill)
        add_text(s,code,x+.1,y+.22,w-.2,.45,size=19,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
        add_text(s,desc,x+.12,y+.9,w-.24,.9,size=12.5,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
        if i<len(stages)-1:
            add_text(s,'→',x+w+.01,y+.85,gap+.2,.45,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'Setiap konfigurasi menambahkan satu perubahan utama secara kumulatif.',1.15,5.15,10.6,.5,size=14,bold=True,align=PP_ALIGN.CENTER)


def slide_researchflow(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1])
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Alur Penelitian',22)
    phases=[
        ('PENGEMBANGAN',['robusta_SNI_Dataset','Split 70 / 15 / 15','Baseline B0','Tetapkan Hard Classes']),
        ('OPTIMASI',['Evaluasi C0–C5','Sensitivity Analysis','Pilih & Bekukan C*']),
        ('KONFIRMASI',['Multi-seed Confirmation','Final Test'])
    ]
    xs=[.55,4.45,8.35]
    for x,(head,items) in zip(xs,phases):
        add_text(s,head,x,1.65,3.1,.45,size=15.5,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
        y=2.25
        for j,it in enumerate(items):
            add_panel(s,x,y,3.1,.72,fill=RGBColor(229,242,220))
            add_text(s,it,x+.1,y+.08,2.9,.5,size=12.8,bold=(it in ('robusta_SNI_Dataset','Pilih & Bekukan C*','Final Test')),align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
            if j<len(items)-1: add_text(s,'↓',x+1.32,y+.72,.45,.27,size=14,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
            y+=.96
    add_text(s,'→',3.75,3.4,.5,.5,size=22,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'→',7.65,3.4,.5,.5,size=22,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_panel(s,2.0,6.0,8.8,.52,fill=RGBColor(217,235,203),radius=False)
    add_text(s,'Test set tidak digunakan untuk memilih C*.',2.2,6.07,8.4,.34,size=13.5,bold=True,align=PP_ALIGN.CENTER)

def slide_cross(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Konfirmasi Lintas Dataset',22)
    add_panel(s,4.3,1.6,4.2,.8,fill=RGBColor(209,232,193)); add_text(s,'C* dibekukan',4.45,1.78,3.9,.45,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    add_text(s,'↓',6.02,2.4,.6,.35,size=20,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
    cards=[('Capstone','14 kelas'),('Lulus','6 kelas'),('Niacubilla','9 kelas')]; xs=[1.0,4.65,8.3]
    for x,(n,c) in zip(xs,cards):
        add_panel(s,x,3.0,3.0,1.7,fill=RGBColor(227,241,217)); add_text(s,n,x+.15,3.25,2.7,.45,size=17,bold=True,align=PP_ALIGN.CENTER); add_text(s,c+'\nB0 vs B3',x+.15,3.8,2.7,.65,size=14,align=PP_ALIGN.CENTER)
    add_text(s,'Seeds: 123 · 2026 · 31415',2.0,5.25,8.8,.42,size=14,bold=True,align=PP_ALIGN.CENTER)
    add_text(s,'Tidak ada retuning C* pada dataset konfirmasi.',2.0,5.75,8.8,.42,size=13,align=PP_ALIGN.CENTER)

def slide_objective(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Tujuan Penelitian',22)
    add_panel(s,.8,2.05,11.0,2.65,fill=PALE_BLUE,radius=False)
    txt='Menganalisis dan mengoptimasi prapemrosesan citra berbasis frekuensi-angular pada YOLO26n untuk deteksi fine-grained cacat biji kopi serta mengevaluasi pengaruhnya terhadap kinerja deteksi dan biaya komputasi.'
    add_text(s,txt,1.25,2.55,10.1,1.65,size=18,align=PP_ALIGN.JUSTIFY,valign=MSO_ANCHOR.MIDDLE)

def slide_eval(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_title(s,'Evaluasi Penelitian',22)
    cols=[('Deteksi',['mAP50–95','mAP50','Precision & Recall']),('Fine-grained',['AP per kelas','AP_H','AP_worst']),('Efisiensi',['Preprocessing time','End-to-end latency','FPS','Peak GPU memory'])]; xs=[.65,4.45,8.25]
    for x,(h,items) in zip(xs,cols):
        add_panel(s,x,1.8,3.25,3.95,fill=PALE_BLUE,radius=False); add_text(s,h,x+.15,2.05,2.95,.45,size=18,bold=True,color=DARK_GREEN,align=PP_ALIGN.CENTER)
        y=2.85
        for it in items: add_text(s,'• '+it,x+.35,y,2.65,.45,size=14); y+=.63

def slide_close(prs,d):
    s=prs.slides.add_slide(prs.slide_layouts[1]);
    for sh in list(s.shapes):
        if getattr(sh,'is_placeholder',False): sh.text_frame.clear()
    add_text(s,'TERIMA KASIH',.9,2.7,7.8,1.0,size=32,bold=True,color=GREEN)
    add_text(s,'Pertanyaan & Diskusi',.95,3.7,5.2,.55,size=16,color=GRAY)

BUILDERS={1:slide_cover,2:slide_agenda,3:slide_intro,4:slide_problem,5:slide_scope,6:slide_related,7:slide_why,8:slide_dataset,9:slide_method,10:slide_preprocess,11:slide_design,12:slide_researchflow,13:slide_cross,14:slide_objective,15:slide_eval,16:slide_close}

def build(template,source,out):
    prs=Presentation(str(template)); delete_all_slides(prs); data=parse_source(source)
    if len(data)!=16: raise ValueError(len(data))
    for d in data: BUILDERS[d['number']](prs,d)
    prs.save(str(out))

if __name__=='__main__':
    ap=argparse.ArgumentParser(); ap.add_argument('--template',required=True); ap.add_argument('--source',required=True); ap.add_argument('--output',required=True); a=ap.parse_args(); build(a.template,a.source,a.output)
